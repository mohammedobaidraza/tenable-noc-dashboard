"""
Tenable Vulnerability Tracker - Interactive NOC Dashboard
==========================================================
Fetches vulnerability data from Tenable.io, updates an Excel tracker,
and generates interactive standalone HTML dashboards with a dark NOC theme.

Usage:
    python tenable_noc_dashboard.py

Requirements:
    pip install plotly pandas openpyxl requests

Output (to C:\\tenable_tracker\\reports\\):
    dashboard.html      - NOC overview with per-tower tabs
    critical.html       - Critical severity breakdown
    high.html           - High severity breakdown
    medium.html         - Medium severity breakdown
    low.html            - Low severity breakdown
    consolidated.html   - All severities combined
"""

import sys
import os
import json
import time
import datetime
import requests
import pandas as pd
import numpy as np
import plotly.graph_objects as go
import plotly.express as px
from plotly.subplots import make_subplots
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, borders
from openpyxl.styles.borders import Border

# ═══════════════════════════════════════════════════════════
# CONFIGURATION
# ═══════════════════════════════════════════════════════════

EXCEL_PATH = r"C:\tenable_tracker\Weekly VM Remediation Dashboard.xlsx"
REPORT_DIR = r"C:\tenable_tracker\reports"
NUM_OF_DAYS = 7
SLA_WEEKS = 6

# Tenable.io API connection parameters
# IMPORTANT: replace these placeholders with your actual Tenable access/secret keys.
accessKey = "YOUR_ACCESS_KEY"
secretKey = "YOUR_SECRET_KEY"

TAGS = [
    {"category": "Tower", "tag": "Intel Server"},
    {"category": "Tower", "tag": "Linux"},
    {"category": "Tower", "tag": "Workstations"},
    {"category": "Tower", "tag": "Network"},
    {"category": "Tower", "tag": "Database"},
    {"category": "Tower", "tag": "CVAN"},
    {"category": "Tower", "tag": "Exchange"},
    {"category": "Tower", "tag": "Middleware"},
    {"category": "Tower", "tag": "Internet(Perimeter) Exposed"},
    {"category": "Tower", "tag": "VMWare"},
    {"category": "[Scan] Cloud", "tag": "GCP"},
    {"category": "Tower", "tag": "EOL"},
]

EXCLUDED_PLUGIN_PREFIXES = ("Linux Distros Unsupported Vulnerability",)

SLA_THRESHOLDS = {"Critical": 30, "High": 60, "Medium": 90, "Low": 180}
AGE_BUCKETS = [(0, 30), (31, 60), (61, 90), (91, 120), (121, None)]
AGE_LABELS = ["0-30 days", "31-60 days", "61-90 days", "91-120 days", "120+ days"]

# NOC Theme
C = {
    "bg": "#0b0b1a",
    "card": "#121230",
    "card_border": "#1e1e4a",
    "accent": "#00d4ff",
    "text": "#e0e0e0",
    "dim": "#8888aa",
    "grid": "#1a1a3a",
    "critical": "#ff2d55",
    "high": "#ff6b35",
    "medium": "#ffbe0b",
    "low": "#06d6a0",
    "info": "#118ab2",
}

SEV_COLORS = {
    "Critical": C["critical"],
    "High": C["high"],
    "Medium": C["medium"],
    "Low": C["low"],
    "Info": C["info"],
}

SEV_ORDER = ["Critical", "High", "Medium", "Low"]


# ═══════════════════════════════════════════════════════════
# TENABLE API
# ═══════════════════════════════════════════════════════════

def is_excluded(record):
    name = (record.get("plugin", {}).get("name", "") or "").strip().lower()
    return any(name.startswith(p.lower()) for p in EXCLUDED_PLUGIN_PREFIXES)


def fetch_tag_data(tag_category, tag, access_key, secret_key):
    """Fetch vulnerability data for a single tag from Tenable.io."""
    print(f"  [{tag_category}:{tag}] Fetching...")

    today = datetime.date.today()
    start_date = today - datetime.timedelta(days=NUM_OF_DAYS)
    epoch = int(time.mktime(start_date.timetuple()))

    headers = {
        "accept": "application/json",
        "content-type": "application/json",
        "X-ApiKeys": f"accessKey={access_key};secretKey={secret_key}",
    }

    payload = {
        "num_assets": 5000,
        "filters": {
            "state": ["OPEN", "REOPENED"],
            "since": epoch,
            f"tag.{tag_category}": [tag],
            "severity": ["critical", "high", "medium", "low"],
        },
    }

    resp = requests.post(
        "https://cloud.tenable.com/vulns/export",
        headers=headers,
        json=payload,
    )
    resp.raise_for_status()
    file_id = resp.json()["export_uuid"]

    while True:
        time.sleep(2)
        status = requests.get(
            f"https://cloud.tenable.com/vulns/export/{file_id}/status",
            headers=headers,
        ).json()
        if status["status"] == "FINISHED":
            break

    download_data = []
    dl_headers = {
        "accept": "application/octet-stream",
        "X-ApiKeys": f"accessKey={access_key};secretKey={secret_key}",
    }
    for chunk_id in status["chunks_available"]:
        chunk = requests.get(
            f"https://cloud.tenable.com/vulns/export/{file_id}/chunks/{chunk_id}",
            headers=dl_headers,
        ).json()
        download_data.extend(chunk)

    print(f"  [{tag}] {len(download_data)} records downloaded.")
    return download_data


def fetch_all_tags(access_key, secret_key):
    """Fetch data for all configured tags. Returns {tag_name: [records]}."""
    all_data = {}
    for t in TAGS:
        try:
            records = fetch_tag_data(t["category"], t["tag"], access_key, secret_key)
            all_data[t["tag"]] = records
        except Exception as e:
            print(f"  [ERROR] {t['tag']}: {e}")
            all_data[t["tag"]] = []
    return all_data


# ═══════════════════════════════════════════════════════════
# DATA PROCESSING
# ═══════════════════════════════════════════════════════════

SEV_MAP = {"critical": "Critical", "high": "High", "medium": "Medium", "low": "Low", "info": "Info"}


def aggregate_by_plugin(records):
    """Aggregate raw Tenable records by plugin ID. Returns dict of plugin details."""
    plugins = {}
    for rec in records:
        if is_excluded(rec):
            continue

        pid = rec["plugin"]["id"]
        sev_raw = rec.get("severity", "info")
        severity = SEV_MAP.get(sev_raw, sev_raw)

        if pid in plugins:
            entry = plugins[pid]
            new_cves = rec["plugin"].get("cve", [])
            entry["cve"] = list(set(entry["cve"]) | set(new_cves))
            entry["count"] += 1
        else:
            ease = rec["plugin"].get("exploitability_ease", "")
            exploit = "Yes" if ease == "Exploits are available" else (
                "No" if ease == "No known exploits are available" else ""
            )

            first_found = rec.get("first_found", "")
            try:
                age = (datetime.datetime.today() - datetime.datetime.strptime(first_found[:10], "%Y-%m-%d")).days
            except (ValueError, IndexError, TypeError):
                age = 0

            plugins[pid] = {
                "cve": rec["plugin"].get("cve", []),
                "cvssv3": rec["plugin"].get("cvss3_base_score", ""),
                "Severity": severity,
                "pluginName": rec["plugin"].get("name", ""),
                "family": rec["plugin"].get("family", ""),
                "exploit": exploit,
                "firstSeen": first_found,
                "age": age,
                "count": 1,
            }

    return plugins


def plugins_to_dataframe(plugins, tower_name=""):
    """Convert plugin dict to a pandas DataFrame."""
    rows = []
    for pid, info in plugins.items():
        rows.append({
            "Tower": tower_name,
            "Plugin ID": pid,
            "CVE": ", ".join(info["cve"]) if info["cve"] else "",
            "CVSS3": info["cvssv3"],
            "Severity": info["Severity"],
            "Plugin Name": info["pluginName"],
            "Family": info["family"],
            "Exploit": info["exploit"],
            "First Seen": info["firstSeen"],
            "Age (Days)": info["age"],
            "Host Count": info["count"],
        })
    df = pd.DataFrame(rows)
    if not df.empty:
        df["SLA Threshold"] = df["Severity"].apply(lambda x: SLA_THRESHOLDS.get(str(x), 9999))
        df["SLA State"] = np.where(df["Age (Days)"] > df["SLA Threshold"], "Past SLA", "Within SLA")
    return df


def compute_age_bucket(age):
    for i, (lo, hi) in enumerate(AGE_BUCKETS):
        if hi is None and age >= lo:
            return AGE_LABELS[i]
        if lo <= age <= hi:
            return AGE_LABELS[i]
    return AGE_LABELS[-1]


# ═══════════════════════════════════════════════════════════
# EXCEL OPERATIONS
# ═══════════════════════════════════════════════════════════

def update_excel(all_plugins, excel_path):
    """Update or create the Excel tracker workbook."""
    os.makedirs(os.path.dirname(excel_path), exist_ok=True)

    if os.path.exists(excel_path):
        wb = openpyxl.load_workbook(excel_path)
    else:
        wb = openpyxl.Workbook()
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

    date_str = datetime.datetime.today().strftime("%m/%d/%y")
    col_name = f"Host Total {date_str}"

    header_font = Font(name="Calibri", size=12, bold=True)
    header_align = Alignment(wrap_text=True, horizontal="center", vertical="center")
    header_fill = PatternFill(start_color="BFBFBF", fill_type="solid")
    center_align = Alignment(wrap_text=False, horizontal="center", vertical="center")
    thin_border_side = borders.Side(style=None, color="000000", border_style="thin")
    thin_border = Border(left=thin_border_side, right=thin_border_side,
                         bottom=thin_border_side, top=thin_border_side)

    HEADERS = ["CVE", "Plugin ID", "CVSS3", "Severity", "Plugin Name",
               "Family", "Exploit", "First Seen", "Age", "SLA State"]

    for tag_name, plugins in all_plugins.items():
        if tag_name not in wb.sheetnames:
            ws = wb.create_sheet(tag_name)
            for i, h in enumerate(HEADERS, 1):
                cell = ws.cell(row=1, column=i, value=h)
                cell.font = header_font
                cell.alignment = header_align
                cell.fill = header_fill
            host_col = len(HEADERS) + 1
        else:
            ws = wb[tag_name]
            ws.insert_cols(15, 1)
            host_col = 15

        hdr_cell = ws.cell(row=1, column=host_col, value=col_name)
        hdr_cell.font = header_font
        hdr_cell.alignment = header_align
        hdr_cell.fill = header_fill

        existing_ids = {}
        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            if row[1].value is not None:
                existing_ids[row[1].value] = row_idx

        new_plugins = []
        for pid, info in plugins.items():
            if pid in existing_ids:
                r = existing_ids[pid]
                ws.cell(row=r, column=1).value = ", ".join(info["cve"]) if info["cve"] else ""
                ws.cell(row=r, column=4).value = info["Severity"]
                ws.cell(row=r, column=7).value = info["exploit"]
                ws.cell(row=r, column=7).alignment = center_align
                ws.cell(row=r, column=8).value = info["firstSeen"]
                ws.cell(row=r, column=9).value = info["age"]
                ws.cell(row=r, column=9).alignment = center_align
                ws.cell(row=r, column=host_col).value = info["count"]
                ws.cell(row=r, column=host_col).alignment = center_align
                ws.cell(row=r, column=host_col).border = thin_border
            else:
                new_plugins.append(pid)

        next_row = ws.max_row + 1
        for pid in new_plugins:
            info = plugins[pid]
            vals = [
                ", ".join(info["cve"]) if info["cve"] else "",
                pid, info["cvssv3"], info["Severity"], info["pluginName"],
                info["family"], info["exploit"], info["firstSeen"], info["age"],
            ]
            for ci, v in enumerate(vals, 1):
                cell = ws.cell(row=next_row, column=ci, value=v)
                cell.alignment = center_align if ci in (2, 3, 4, 7, 8, 9) else Alignment(
                    wrap_text=False, horizontal="left", vertical="center")
                cell.border = thin_border

            sla_formula = (
                f'=IF(D{next_row}="Critical",IF(I{next_row}>30,"Past SLA","Within SLA"),'
                f'IF(D{next_row}="High",IF(I{next_row}>60,"Past SLA","Within SLA"),'
                f'IF(D{next_row}="Medium",IF(I{next_row}>90,"Past SLA","Within SLA"),'
                f'IF(D{next_row}="Low",IF(I{next_row}>180,"Past SLA","Within SLA"),"N/A"))))'
            )
            ws.cell(row=next_row, column=10, value=sla_formula).alignment = center_align
            ws.cell(row=next_row, column=10).border = thin_border

            ws.cell(row=next_row, column=host_col, value=info["count"])
            ws.cell(row=next_row, column=host_col).alignment = center_align
            ws.cell(row=next_row, column=host_col).border = thin_border

            next_row += 1

        print(f"  [{tag_name}] Excel updated: {len(new_plugins)} new, {len(existing_ids)} existing.")

    wb.save(excel_path)
    print(f"  Excel saved: {excel_path}")


def load_excel_to_dataframes(excel_path):
    """Read the Excel workbook and return {tag_name: DataFrame} + weekly trend info."""
    if not os.path.exists(excel_path):
        return {}, {}

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    tag_dfs = {}
    weekly_info = {}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            continue

        headers = [str(h) if h else f"col_{i}" for i, h in enumerate(rows[0])]
        data = rows[1:]
        df = pd.DataFrame(data, columns=headers)
        df = df.dropna(how="all")

        host_cols = [h for h in headers if str(h).startswith("Host Total")]
        week_dates = []
        for hc in host_cols:
            try:
                ds = hc.replace("Host Total ", "").strip()
                dt = pd.to_datetime(ds, format="%m/%d/%y")
                week_dates.append({"date": dt, "col": hc})
            except (ValueError, TypeError):
                pass
        week_dates.sort(key=lambda x: x["date"])

        std_cols = {
            "CVE": None, "Plugin ID": None, "CVSS3": None, "Severity": None,
            "Plugin Name": None, "Family": None, "Exploit": None,
            "First Seen": None, "Age": None, "SLA State": None,
        }
        for target in std_cols:
            for h in headers:
                if target.lower() in h.lower():
                    std_cols[target] = h
                    break

        rename_map = {v: k for k, v in std_cols.items() if v and v != k}
        df = df.rename(columns=rename_map)

        if "Age" in df.columns:
            df["Age (Days)"] = pd.to_numeric(df["Age"], errors="coerce").fillna(0).astype(int)
        elif "Age (Days)" not in df.columns:
            df["Age (Days)"] = 0

        latest_host_col = week_dates[-1]["col"] if week_dates else None
        if latest_host_col and latest_host_col in df.columns:
            df["Host Count"] = pd.to_numeric(df[latest_host_col], errors="coerce").fillna(0).astype(int)
        elif "Host Count" not in df.columns:
            df["Host Count"] = 0

        if "Severity" not in df.columns:
            df["Severity"] = "Unknown"

        df["Tower"] = sheet_name

        if "SLA State" not in df.columns:
            df["SLA Threshold"] = df["Severity"].apply(lambda x: SLA_THRESHOLDS.get(str(x), 9999))
            df["SLA State"] = np.where(df["Age (Days)"] > df["SLA Threshold"], "Past SLA", "Within SLA")

        tag_dfs[sheet_name] = df
        weekly_info[sheet_name] = week_dates[-SLA_WEEKS:] if len(week_dates) >= SLA_WEEKS else week_dates

    wb.close()
    return tag_dfs, weekly_info


# ═══════════════════════════════════════════════════════════
# PLOTLY CHART BUILDERS
# ═══════════════════════════════════════════════════════════

def _base_layout(**overrides):
    layout = dict(
        paper_bgcolor=C["bg"],
        plot_bgcolor=C["card"],
        font=dict(color=C["text"], family="Segoe UI, Consolas, monospace", size=12),
        margin=dict(l=50, r=30, t=50, b=40),
        xaxis=dict(gridcolor=C["grid"], zerolinecolor=C["grid"]),
        yaxis=dict(gridcolor=C["grid"], zerolinecolor=C["grid"]),
        legend=dict(bgcolor="rgba(0,0,0,0)", font=dict(color=C["text"])),
    )
    layout.update(overrides)
    return layout


def chart_severity_pie(df, title="Severity Distribution"):
    counts = df.groupby("Severity")["Host Count"].sum().reindex(SEV_ORDER, fill_value=0)
    fig = go.Figure(go.Pie(
        labels=counts.index.tolist(),
        values=counts.values.tolist(),
        marker=dict(colors=[SEV_COLORS.get(s, C["info"]) for s in counts.index]),
        hole=0.45,
        textinfo="label+percent+value",
        textfont=dict(size=13, color="#fff"),
    ))
    fig.update_layout(**_base_layout(title=dict(text=title, font=dict(size=16, color=C["accent"]))))
    return fig


def chart_severity_bar(df, title="Severity Distribution"):
    counts = df.groupby("Severity")["Host Count"].sum().reindex(SEV_ORDER, fill_value=0)
    fig = go.Figure(go.Bar(
        x=counts.index.tolist(),
        y=counts.values.tolist(),
        marker_color=[SEV_COLORS.get(s, C["info"]) for s in counts.index],
        text=counts.values.tolist(),
        textposition="outside",
        textfont=dict(color=C["text"]),
    ))
    fig.update_layout(**_base_layout(title=dict(text=title, font=dict(size=16, color=C["accent"]))))
    return fig


def chart_top_vulns(df, n=10, title="Top Vulnerabilities by Host Count"):
    if df.empty:
        return go.Figure().update_layout(**_base_layout(title=dict(text=title)))

    top = df.nlargest(n, "Host Count")[["Plugin Name", "Host Count", "Severity"]].copy()
    top = top.iloc[::-1]
    colors = [SEV_COLORS.get(s, C["info"]) for s in top["Severity"]]

    fig = go.Figure(go.Bar(
        y=top["Plugin Name"].tolist(),
        x=top["Host Count"].tolist(),
        orientation="h",
        marker_color=colors,
        text=top["Host Count"].tolist(),
        textposition="outside",
        textfont=dict(color=C["text"]),
    ))
    fig.update_layout(**_base_layout(
        title=dict(text=title, font=dict(size=16, color=C["accent"])),
        yaxis=dict(gridcolor=C["grid"], tickfont=dict(size=10)),
        height=max(350, n * 35 + 100),
    ))
    return fig


def chart_age_histogram(df, title="Vulnerability Age Distribution"):
    if df.empty or "Age (Days)" not in df.columns:
        return go.Figure().update_layout(**_base_layout(title=dict(text=title)))

    fig = go.Figure()
    for sev in SEV_ORDER:
        subset = df[df["Severity"] == sev]
        if not subset.empty:
            fig.add_trace(go.Histogram(
                x=subset["Age (Days)"].tolist(),
                name=sev,
                marker_color=SEV_COLORS.get(sev, C["info"]),
                opacity=0.8,
            ))
    fig.update_layout(**_base_layout(
        title=dict(text=title, font=dict(size=16, color=C["accent"])),
        barmode="stack",
        xaxis_title="Age (Days)",
        yaxis_title="Count",
    ))
    return fig


def chart_exploit_pie(df, title="Exploit Availability"):
    if "Exploit" not in df.columns:
        return go.Figure().update_layout(**_base_layout(title=dict(text=title)))

    counts = df["Exploit"].value_counts()
    color_map = {"Yes": C["critical"], "No": C["low"], "": C["dim"]}
    fig = go.Figure(go.Pie(
        labels=counts.index.tolist(),
        values=counts.values.tolist(),
        marker=dict(colors=[color_map.get(l, C["dim"]) for l in counts.index]),
        hole=0.4,
        textinfo="label+percent+value",
        textfont=dict(color="#fff"),
    ))
    fig.update_layout(**_base_layout(title=dict(text=title, font=dict(size=16, color=C["accent"]))))
    return fig


def chart_family_bar(df, n=12, title="Plugin Family Breakdown"):
    if "Family" not in df.columns or df.empty:
        return go.Figure().update_layout(**_base_layout(title=dict(text=title)))

    fam = df.groupby("Family")["Host Count"].sum().nlargest(n).sort_values()
    fig = go.Figure(go.Bar(
        y=fam.index.tolist(),
        x=fam.values.tolist(),
        orientation="h",
        marker_color=C["accent"],
        text=fam.values.tolist(),
        textposition="outside",
        textfont=dict(color=C["text"]),
    ))
    fig.update_layout(**_base_layout(
        title=dict(text=title, font=dict(size=16, color=C["accent"])),
        height=max(350, n * 30 + 100),
    ))
    return fig


def chart_sla_compliance(df, title="SLA Compliance"):
    if "SLA State" not in df.columns or df.empty:
        return go.Figure().update_layout(**_base_layout(title=dict(text=title)))

    total = len(df)
    within = (df["SLA State"] == "Within SLA").sum()
    past = total - within
    pct = round(within / total * 100, 1) if total > 0 else 0

    fig = go.Figure(go.Pie(
        labels=["Within SLA", "Past SLA"],
        values=[within, past],
        marker=dict(colors=[C["low"], C["critical"]]),
        hole=0.6,
        textinfo="label+percent",
        textfont=dict(color="#fff", size=13),
    ))
    fig.add_annotation(
        text=f"<b>{pct}%</b>", x=0.5, y=0.5, font=dict(size=28, color=C["accent"]),
        showarrow=False, xref="paper", yref="paper",
    )
    fig.update_layout(**_base_layout(title=dict(text=title, font=dict(size=16, color=C["accent"]))))
    return fig


def chart_aging_buckets(df, title="Aging Buckets"):
    if df.empty:
        return go.Figure().update_layout(**_base_layout(title=dict(text=title)))

    df = df.copy()
    df["Bucket"] = df["Age (Days)"].apply(compute_age_bucket)

    fig = go.Figure()
    for sev in SEV_ORDER:
        subset = df[df["Severity"] == sev]
        if subset.empty:
            continue
        bucket_counts = subset.groupby("Bucket").size().reindex(AGE_LABELS, fill_value=0)
        fig.add_trace(go.Bar(
            x=AGE_LABELS,
            y=bucket_counts.values.tolist(),
            name=sev,
            marker_color=SEV_COLORS.get(sev, C["info"]),
        ))
    fig.update_layout(**_base_layout(
        title=dict(text=title, font=dict(size=16, color=C["accent"])),
        barmode="group",
        xaxis_title="Age Bucket",
        yaxis_title="Vulnerability Count",
    ))
    return fig


def chart_tower_comparison(tag_dfs, title="Tower Comparison by Severity"):
    towers, crits, highs, meds, lows = [], [], [], [], []
    for tname, df in tag_dfs.items():
        towers.append(tname)
        sev_counts = df.groupby("Severity")["Host Count"].sum()
        crits.append(int(sev_counts.get("Critical", 0)))
        highs.append(int(sev_counts.get("High", 0)))
        meds.append(int(sev_counts.get("Medium", 0)))
        lows.append(int(sev_counts.get("Low", 0)))

    fig = go.Figure()
    for name, vals, color in [
        ("Critical", crits, C["critical"]),
        ("High", highs, C["high"]),
        ("Medium", meds, C["medium"]),
        ("Low", lows, C["low"]),
    ]:
        fig.add_trace(go.Bar(x=towers, y=vals, name=name, marker_color=color))

    fig.update_layout(**_base_layout(
        title=dict(text=title, font=dict(size=16, color=C["accent"])),
        barmode="stack",
        xaxis=dict(tickangle=-45, gridcolor=C["grid"]),
        height=500,
    ))
    return fig


def chart_treemap(combined_df, title="Vulnerability Map"):
    if combined_df.empty:
        return go.Figure().update_layout(**_base_layout(title=dict(text=title)))

    grouped = combined_df.groupby(["Tower", "Severity", "Family"])["Host Count"].sum().reset_index()
    grouped = grouped[grouped["Host Count"] > 0]

    if grouped.empty:
        return go.Figure().update_layout(**_base_layout(title=dict(text=title)))

    color_discrete = {s: SEV_COLORS.get(s, C["info"]) for s in SEV_ORDER}

    fig = px.treemap(
        grouped,
        path=["Tower", "Severity", "Family"],
        values="Host Count",
        color="Severity",
        color_discrete_map=color_discrete,
        title=title,
    )
    fig.update_layout(
        paper_bgcolor=C["bg"],
        plot_bgcolor=C["card"],
        font=dict(color=C["text"], family="Segoe UI, Consolas, monospace"),
        margin=dict(l=10, r=10, t=50, b=10),
        height=550,
    )
    fig.update_traces(
        textfont=dict(size=12, color="#fff"),
        marker=dict(cornerradius=4),
    )
    return fig


def chart_fluctuation_heatmap(tag_dfs, weekly_info, title="Weekly Fluctuation Heatmap"):
    """Heatmap: rows=towers, cols=weeks, color=total host count change %."""
    towers_with_data = []
    all_weeks_set = set()
    tower_week_totals = {}

    for tname, df in tag_dfs.items():
        weeks = weekly_info.get(tname, [])
        if len(weeks) < 2:
            continue
        towers_with_data.append(tname)
        week_totals = {}
        for w in weeks:
            col = w["col"]
            if col in df.columns:
                total = pd.to_numeric(df[col], errors="coerce").sum()
                date_label = w["date"].strftime("%m/%d")
                week_totals[date_label] = total
                all_weeks_set.add(date_label)
        tower_week_totals[tname] = week_totals

    if not towers_with_data or len(all_weeks_set) < 2:
        fig = go.Figure()
        fig.add_annotation(
            text="Insufficient historical data for fluctuation analysis.<br>Run weekly to build trend data.",
            x=0.5, y=0.5, showarrow=False, font=dict(color=C["dim"], size=14),
            xref="paper", yref="paper",
        )
        fig.update_layout(**_base_layout(title=dict(text=title, font=dict(size=16, color=C["accent"])), height=300))
        return fig

    week_labels = sorted(all_weeks_set)
    z_data = []
    for tname in towers_with_data:
        row = []
        totals = tower_week_totals.get(tname, {})
        prev = None
        for wl in week_labels:
            val = totals.get(wl, 0)
            if prev is not None and prev > 0:
                change_pct = ((val - prev) / prev) * 100
            else:
                change_pct = 0
            row.append(round(change_pct, 1))
            prev = val
        z_data.append(row)

    fig = go.Figure(go.Heatmap(
        z=z_data,
        x=week_labels,
        y=towers_with_data,
        colorscale=[[0, C["low"]], [0.5, C["bg"]], [1, C["critical"]]],
        zmid=0,
        text=[[f"{v:+.1f}%" for v in row] for row in z_data],
        texttemplate="%{text}",
        textfont=dict(size=11, color="#fff"),
        colorbar=dict(title="% Change", ticksuffix="%", bgcolor=C["bg"], tickfont=dict(color=C["text"])),
    ))
    fig.update_layout(**_base_layout(
        title=dict(text=title, font=dict(size=16, color=C["accent"])),
        height=max(300, len(towers_with_data) * 40 + 120),
        xaxis_title="Week",
        yaxis=dict(gridcolor=C["grid"], autorange="reversed"),
    ))
    return fig


def chart_weekly_trend(df, weekly_info_list, title="Weekly Host Count Trend"):
    """Line chart of total host counts per week for a single tower."""
    if not weekly_info_list or df.empty:
        fig = go.Figure()
        fig.add_annotation(
            text="Run the tracker weekly to build trend data.",
            x=0.5, y=0.5, showarrow=False, font=dict(color=C["dim"], size=14),
            xref="paper", yref="paper",
        )
        fig.update_layout(**_base_layout(title=dict(text=title, font=dict(size=16, color=C["accent"])), height=300))
        return fig

    dates, totals = [], []
    for w in weekly_info_list:
        col = w["col"]
        if col in df.columns:
            dates.append(w["date"])
            totals.append(pd.to_numeric(df[col], errors="coerce").sum())

    fig = go.Figure(go.Scatter(
        x=dates,
        y=totals,
        mode="lines+markers+text",
        line=dict(color=C["accent"], width=3),
        marker=dict(size=10, color=C["accent"], line=dict(color="#fff", width=1)),
        text=[str(int(t)) for t in totals],
        textposition="top center",
        textfont=dict(color=C["text"], size=12),
        fill="tozeroy",
        fillcolor="rgba(0,212,255,0.08)",
    ))
    fig.update_layout(**_base_layout(
        title=dict(text=title, font=dict(size=16, color=C["accent"])),
        xaxis_title="Week",
        yaxis_title="Total Hosts Affected",
        height=350,
    ))
    return fig


def chart_sla_per_tower(tag_dfs, title="SLA Compliance by Tower"):
    towers, pcts = [], []
    for tname, df in tag_dfs.items():
        if df.empty or "SLA State" not in df.columns:
            continue
        total = len(df)
        within = (df["SLA State"] == "Within SLA").sum()
        pcts.append(round(within / total * 100, 1) if total > 0 else 0)
        towers.append(tname)

    colors = [C["low"] if p >= 80 else C["medium"] if p >= 50 else C["critical"] for p in pcts]

    fig = go.Figure(go.Bar(
        x=towers,
        y=pcts,
        marker_color=colors,
        text=[f"{p}%" for p in pcts],
        textposition="outside",
        textfont=dict(color=C["text"]),
    ))
    fig.update_layout(**_base_layout(
        title=dict(text=title, font=dict(size=16, color=C["accent"])),
        yaxis=dict(range=[0, 110], gridcolor=C["grid"]),
        xaxis=dict(tickangle=-45),
    ))
    fig.add_hline(y=80, line_dash="dash", line_color=C["low"], annotation_text="80% target",
                  annotation_font_color=C["low"])
    return fig


# ═══════════════════════════════════════════════════════════
# HTML GENERATION
# ═══════════════════════════════════════════════════════════

NOC_CSS = """
<style>
    * { box-sizing: border-box; margin: 0; padding: 0; }
    body {
        background: %(bg)s; color: %(text)s;
        font-family: 'Segoe UI', 'Consolas', monospace;
        padding: 0; margin: 0;
    }
    .header {
        background: linear-gradient(135deg, %(card)s 0%%, #0d0d2b 100%%);
        border-bottom: 2px solid %(accent)s;
        padding: 24px 40px;
        display: flex; align-items: center; justify-content: space-between;
    }
    .header h1 {
        font-size: 1.6rem; color: %(accent)s; font-weight: 700;
        text-transform: uppercase; letter-spacing: 2px;
    }
    .header .timestamp {
        color: %(dim)s; font-size: 0.85rem;
    }
    .kpi-row {
        display: flex; gap: 16px; padding: 24px 40px; flex-wrap: wrap;
    }
    .kpi-card {
        background: %(card)s; border: 1px solid %(card_border)s;
        border-radius: 10px; padding: 20px 28px; text-align: center;
        flex: 1; min-width: 140px;
        box-shadow: 0 0 20px rgba(0,212,255,0.04);
        transition: box-shadow 0.3s;
    }
    .kpi-card:hover { box-shadow: 0 0 30px rgba(0,212,255,0.12); }
    .kpi-value {
        font-size: 2.2rem; font-weight: 800;
        font-family: 'Consolas', monospace;
    }
    .kpi-label {
        font-size: 0.75rem; color: %(dim)s;
        text-transform: uppercase; letter-spacing: 1.5px; margin-top: 6px;
    }
    .chart-grid {
        display: grid;
        grid-template-columns: repeat(auto-fit, minmax(520px, 1fr));
        gap: 20px; padding: 0 40px 24px;
    }
    .chart-card {
        background: %(card)s; border: 1px solid %(card_border)s;
        border-radius: 10px; padding: 12px; overflow: hidden;
        box-shadow: 0 0 15px rgba(0,0,0,0.3);
    }
    .chart-full { grid-column: 1 / -1; }
    .tabs {
        display: flex; gap: 4px; padding: 0 40px; flex-wrap: wrap;
        border-bottom: 1px solid %(card_border)s;
    }
    .tab-btn {
        background: %(card)s; color: %(dim)s; border: 1px solid %(card_border)s;
        border-bottom: none; padding: 10px 20px; cursor: pointer;
        border-radius: 8px 8px 0 0; font-size: 0.85rem;
        font-family: 'Segoe UI', monospace; transition: all 0.2s;
    }
    .tab-btn:hover { color: %(text)s; background: #1a1a3e; }
    .tab-btn.active {
        color: %(accent)s; background: %(bg)s;
        border-color: %(accent)s; border-bottom: 2px solid %(bg)s;
        font-weight: 700;
    }
    .tab-content { display: none; padding: 24px 0; }
    .tab-content.active { display: block; }
    .section-title {
        font-size: 1.1rem; color: %(accent)s; margin: 24px 40px 12px;
        text-transform: uppercase; letter-spacing: 1px;
        border-left: 3px solid %(accent)s; padding-left: 12px;
    }
    table.data-table {
        width: calc(100%% - 80px); margin: 16px 40px;
        border-collapse: collapse; font-size: 0.8rem;
    }
    table.data-table th {
        background: %(card)s; color: %(accent)s; padding: 10px 12px;
        border: 1px solid %(card_border)s; text-align: left;
        text-transform: uppercase; letter-spacing: 0.5px; font-size: 0.7rem;
        position: sticky; top: 0; cursor: pointer;
    }
    table.data-table th:hover { background: #1a1a3e; }
    table.data-table td {
        padding: 8px 12px; border: 1px solid %(card_border)s;
    }
    table.data-table tr:nth-child(even) { background: rgba(255,255,255,0.02); }
    table.data-table tr:hover { background: rgba(0,212,255,0.06); }
    .sev-critical { color: %(critical)s; font-weight: 700; }
    .sev-high { color: %(high)s; font-weight: 700; }
    .sev-medium { color: %(medium)s; font-weight: 700; }
    .sev-low { color: %(low)s; font-weight: 700; }
    .sla-past { color: %(critical)s; font-weight: 700; }
    .sla-within { color: %(low)s; }
    .chart-toggle {
        display: flex; gap: 4px; justify-content: flex-end;
        padding: 4px 8px;
    }
    .toggle-btn {
        background: %(card_border)s; color: %(dim)s; border: none;
        padding: 4px 12px; border-radius: 4px; cursor: pointer;
        font-size: 0.7rem; font-family: monospace; transition: all 0.2s;
    }
    .toggle-btn.active { background: %(accent)s; color: #000; font-weight: 700; }
    .toggle-btn:hover { background: #2a2a5a; color: %(text)s; }
    .nav-links {
        display: flex; gap: 10px; padding: 16px 40px; flex-wrap: wrap;
    }
    .nav-link {
        background: %(card)s; color: %(accent)s; border: 1px solid %(card_border)s;
        padding: 8px 18px; border-radius: 6px; text-decoration: none;
        font-size: 0.8rem; transition: all 0.2s;
    }
    .nav-link:hover { background: #1a1a3e; border-color: %(accent)s; }
    .nav-link.active { background: %(accent)s; color: #000; font-weight: 700; }
    .search-box {
        background: %(card)s; border: 1px solid %(card_border)s; color: %(text)s;
        padding: 8px 16px; border-radius: 6px; font-size: 0.85rem;
        width: 300px; margin: 0 40px 16px;
        font-family: monospace;
    }
    .search-box:focus { outline: none; border-color: %(accent)s; }
    @media (max-width: 768px) {
        .chart-grid { grid-template-columns: 1fr; padding: 0 16px 16px; }
        .kpi-row { padding: 16px; }
        .header { padding: 16px; }
        table.data-table { width: calc(100%% - 32px); margin: 8px 16px; }
    }
</style>
""" % C

SORT_JS = """
<script>
function sortTable(tableId, colIdx) {
    const table = document.getElementById(tableId);
    const tbody = table.querySelector('tbody') || table;
    const rows = Array.from(tbody.querySelectorAll('tr'));
    const header = table.querySelectorAll('th')[colIdx];
    const asc = header.dataset.sort !== 'asc';
    header.dataset.sort = asc ? 'asc' : 'desc';

    rows.sort((a, b) => {
        let va = a.cells[colIdx]?.textContent.trim() || '';
        let vb = b.cells[colIdx]?.textContent.trim() || '';
        let na = parseFloat(va), nb = parseFloat(vb);
        if (!isNaN(na) && !isNaN(nb)) return asc ? na - nb : nb - na;
        return asc ? va.localeCompare(vb) : vb.localeCompare(va);
    });
    rows.forEach(r => tbody.appendChild(r));
}
function filterTable(tableId, query) {
    const table = document.getElementById(tableId);
    const rows = table.querySelectorAll('tbody tr, tr:not(:first-child)');
    const q = query.toLowerCase();
    rows.forEach(r => {
        r.style.display = r.textContent.toLowerCase().includes(q) ? '' : 'none';
    });
}
function showTab(tabGroup, tabId) {
    document.querySelectorAll('[data-tab-group="'+tabGroup+'"]').forEach(el => {
        el.classList.remove('active');
    });
    document.querySelectorAll('[data-tab-id="'+tabId+'"]').forEach(el => {
        el.classList.add('active');
    });
}
function toggleChart(containerId, chartType) {
    const container = document.getElementById(containerId);
    container.querySelectorAll('.chart-variant').forEach(el => {
        el.style.display = 'none';
    });
    const target = container.querySelector('[data-chart-type="'+chartType+'"]');
    if (target) target.style.display = 'block';
    container.querySelectorAll('.toggle-btn').forEach(btn => {
        btn.classList.toggle('active', btn.dataset.chartType === chartType);
    });
}
</script>
"""


def _fig_to_div(fig, height=None):
    if height:
        fig.update_layout(height=height)
    return fig.to_html(include_plotlyjs=False, full_html=False, config={"displayModeBar": True, "responsive": True})


def _sev_class(sev):
    return f"sev-{sev.lower()}" if sev in SEV_ORDER else ""


def _sla_class(state):
    return "sla-past" if state == "Past SLA" else "sla-within"


def _build_data_table(df, table_id="data-table"):
    display_cols = ["Tower", "Plugin ID", "CVE", "CVSS3", "Severity",
                    "Plugin Name", "Family", "Exploit", "Age (Days)", "Host Count", "SLA State"]
    cols = [c for c in display_cols if c in df.columns]
    if not cols:
        return "<p style='color:#8888aa;padding:40px;'>No data available.</p>"

    html_parts = [
        f'<input type="text" class="search-box" placeholder="Search vulnerabilities..." '
        f'oninput="filterTable(\'{table_id}\', this.value)">',
        f'<table class="data-table" id="{table_id}">',
        "<thead><tr>",
    ]
    for i, col in enumerate(cols):
        html_parts.append(f'<th onclick="sortTable(\'{table_id}\', {i})">{col} &#x25B4;&#x25BE;</th>')
    html_parts.append("</tr></thead><tbody>")

    for _, row in df[cols].iterrows():
        html_parts.append("<tr>")
        for col in cols:
            val = row[col] if pd.notna(row[col]) else ""
            td_class = ""
            if col == "Severity":
                td_class = f' class="{_sev_class(str(val))}"'
            elif col == "SLA State":
                td_class = f' class="{_sla_class(str(val))}"'
            html_parts.append(f"<td{td_class}>{val}</td>")
        html_parts.append("</tr>")

    html_parts.append("</tbody></table>")
    return "\n".join(html_parts)


def _switchable_chart(container_id, charts_dict, default="pie"):
    """Build a chart card with toggle buttons for switching chart types."""
    toggle_btns = []
    chart_divs = []

    for chart_type, fig in charts_dict.items():
        is_default = chart_type == default
        active_cls = " active" if is_default else ""
        display = "block" if is_default else "none"

        toggle_btns.append(
            f'<button class="toggle-btn{active_cls}" data-chart-type="{chart_type}" '
            f'onclick="toggleChart(\'{container_id}\', \'{chart_type}\')">'
            f'{chart_type.title()}</button>'
        )
        chart_divs.append(
            f'<div class="chart-variant" data-chart-type="{chart_type}" style="display:{display}">'
            f'{_fig_to_div(fig)}</div>'
        )

    return (
        f'<div class="chart-card" id="{container_id}">'
        f'<div class="chart-toggle">{"".join(toggle_btns)}</div>'
        f'{"".join(chart_divs)}</div>'
    )


def _kpi_card(label, value, color=None):
    color = color or C["accent"]
    return (
        f'<div class="kpi-card">'
        f'<div class="kpi-value" style="color:{color}">{value}</div>'
        f'<div class="kpi-label">{label}</div>'
        f'</div>'
    )


def _page_wrapper(title, body_html, nav_links=None):
    nav = ""
    if nav_links:
        links_html = "".join(
            f'<a href="{href}" class="nav-link{" active" if active else ""}">{label}</a>'
            for label, href, active in nav_links
        )
        nav = f'<div class="nav-links">{links_html}</div>'

    timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{title}</title>
<script src="https://cdn.plot.ly/plotly-2.35.2.min.js"></script>
{NOC_CSS}
{SORT_JS}
</head>
<body>
<div class="header">
    <h1>{title}</h1>
    <span class="timestamp">Generated: {timestamp}</span>
</div>
{nav}
{body_html}
</body>
</html>"""


def _nav_links(active_page="dashboard"):
    pages = [
        ("Dashboard", "dashboard.html", "dashboard"),
        ("Critical", "critical.html", "critical"),
        ("High", "high.html", "high"),
        ("Medium", "medium.html", "medium"),
        ("Low", "low.html", "low"),
        ("Consolidated", "consolidated.html", "consolidated"),
    ]
    return [(label, href, page_id == active_page) for label, href, page_id in pages]


# ─── Generate: severity HTML ─────────────────────────────

def generate_severity_html(combined_df, severity, output_path):
    """Generate a standalone HTML report for a single severity level."""
    if severity == "All":
        filtered = combined_df.copy()
        title = "Consolidated Vulnerability Report"
        page_id = "consolidated"
    else:
        filtered = combined_df[combined_df["Severity"] == severity].copy()
        title = f"{severity} Vulnerabilities Report"
        page_id = severity.lower()

    total = len(filtered)
    hosts = int(filtered["Host Count"].sum()) if not filtered.empty else 0
    exploitable = int((filtered.get("Exploit") == "Yes").sum()) if "Exploit" in filtered.columns else 0
    avg_age = int(filtered["Age (Days)"].mean()) if not filtered.empty else 0

    if "SLA State" in filtered.columns and total > 0:
        within = int((filtered["SLA State"] == "Within SLA").sum())
        sla_pct = round(within / total * 100, 1)
    else:
        sla_pct = 0

    kpi_html = '<div class="kpi-row">'
    kpi_html += _kpi_card("Total Vulns", total, C.get(severity.lower(), C["accent"]))
    kpi_html += _kpi_card("Total Hosts", hosts, C["accent"])
    kpi_html += _kpi_card("Exploitable", exploitable, C["critical"])
    kpi_html += _kpi_card("Avg Age (Days)", avg_age, C["medium"])
    kpi_html += _kpi_card("SLA Compliance", f"{sla_pct}%", C["low"] if sla_pct >= 80 else C["critical"])
    kpi_html += "</div>"

    charts_html = '<div class="chart-grid">'

    if severity == "All":
        charts_html += _switchable_chart("sev-dist", {
            "pie": chart_severity_pie(filtered),
            "bar": chart_severity_bar(filtered),
        })
    else:
        charts_html += f'<div class="chart-card">{_fig_to_div(chart_sla_compliance(filtered))}</div>'

    charts_html += f'<div class="chart-card">{_fig_to_div(chart_age_histogram(filtered))}</div>'

    tower_counts = filtered.groupby("Tower")["Host Count"].sum().sort_values(ascending=True)
    if not tower_counts.empty:
        fig_tower = go.Figure(go.Bar(
            y=tower_counts.index.tolist(), x=tower_counts.values.tolist(),
            orientation="h",
            marker_color=C.get(severity.lower(), C["accent"]),
            text=tower_counts.values.tolist(), textposition="outside",
            textfont=dict(color=C["text"]),
        ))
        fig_tower.update_layout(**_base_layout(
            title=dict(text="Distribution by Tower", font=dict(size=16, color=C["accent"])),
            height=max(300, len(tower_counts) * 30 + 100),
        ))
        charts_html += f'<div class="chart-card">{_fig_to_div(fig_tower)}</div>'

    charts_html += _switchable_chart(f"family-{page_id}", {
        "bar": chart_family_bar(filtered, title="Plugin Family Breakdown"),
        "pie": chart_exploit_pie(filtered, title="Exploit Availability"),
    }, default="bar")

    charts_html += f'<div class="chart-card chart-full">{_fig_to_div(chart_top_vulns(filtered, n=15))}</div>'

    charts_html += f'<div class="chart-card chart-full">{_fig_to_div(chart_aging_buckets(filtered))}</div>'

    charts_html += "</div>"

    table_html = _build_data_table(filtered, f"tbl-{page_id}")

    body = kpi_html + charts_html + '<div class="section-title">Vulnerability Details</div>' + table_html

    full_html = _page_wrapper(title, body, _nav_links(page_id))

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(full_html)
    print(f"  Generated: {output_path}")


# ─── Generate: main dashboard HTML ───────────────────────

def generate_dashboard_html(tag_dfs, weekly_info, output_path):
    """Generate the main NOC dashboard with per-tower tabs."""
    combined = pd.concat(tag_dfs.values(), ignore_index=True) if tag_dfs else pd.DataFrame()

    total = len(combined)
    hosts_total = int(combined["Host Count"].sum()) if not combined.empty else 0
    crit_count = int(combined[combined["Severity"] == "Critical"]["Host Count"].sum()) if not combined.empty else 0
    high_count = int(combined[combined["Severity"] == "High"]["Host Count"].sum()) if not combined.empty else 0
    med_count = int(combined[combined["Severity"] == "Medium"]["Host Count"].sum()) if not combined.empty else 0
    low_count = int(combined[combined["Severity"] == "Low"]["Host Count"].sum()) if not combined.empty else 0
    exploitable = int((combined.get("Exploit") == "Yes").sum()) if "Exploit" in combined.columns else 0

    # ── Overview KPIs ──
    kpi = '<div class="kpi-row">'
    kpi += _kpi_card("Total Vulns", total, C["accent"])
    kpi += _kpi_card("Total Hosts", hosts_total, C["accent"])
    kpi += _kpi_card("Critical", crit_count, C["critical"])
    kpi += _kpi_card("High", high_count, C["high"])
    kpi += _kpi_card("Medium", med_count, C["medium"])
    kpi += _kpi_card("Low", low_count, C["low"])
    kpi += _kpi_card("Exploitable", exploitable, "#ff4757")
    kpi += "</div>"

    # ── Overview charts ──
    overview_charts = '<div class="chart-grid">'
    overview_charts += _switchable_chart("overview-sev", {
        "pie": chart_severity_pie(combined, "Severity Distribution"),
        "bar": chart_severity_bar(combined, "Severity Distribution"),
    })
    overview_charts += f'<div class="chart-card">{_fig_to_div(chart_sla_per_tower(tag_dfs))}</div>'
    overview_charts += f'<div class="chart-card chart-full">{_fig_to_div(chart_tower_comparison(tag_dfs))}</div>'
    overview_charts += f'<div class="chart-card chart-full">{_fig_to_div(chart_treemap(combined, "Vulnerability Map - Tower / Severity / Family"))}</div>'
    overview_charts += f'<div class="chart-card chart-full">{_fig_to_div(chart_fluctuation_heatmap(tag_dfs, weekly_info))}</div>'
    overview_charts += f'<div class="chart-card chart-full">{_fig_to_div(chart_aging_buckets(combined, "Overall Aging Buckets"))}</div>'
    overview_charts += "</div>"

    # ── Tower tabs ──
    tab_buttons = '<div class="tabs">'
    tab_buttons += '<button class="tab-btn active" data-tab-group="tower" data-tab-id="overview" onclick="showTab(\'tower\',\'overview\')">Overview</button>'

    tower_tab_contents = ""

    for i, tname in enumerate(tag_dfs.keys()):
        safe_id = tname.replace(" ", "_").replace("(", "").replace(")", "")
        tab_buttons += (
            f'<button class="tab-btn" data-tab-group="tower" data-tab-id="{safe_id}" '
            f'onclick="showTab(\'tower\',\'{safe_id}\')">{tname}</button>'
        )

        tdf = tag_dfs[tname]
        tw_info = weekly_info.get(tname, [])

        t_total = len(tdf)
        t_hosts = int(tdf["Host Count"].sum()) if not tdf.empty else 0
        t_crit = int(tdf[tdf["Severity"] == "Critical"]["Host Count"].sum()) if not tdf.empty else 0
        t_high = int(tdf[tdf["Severity"] == "High"]["Host Count"].sum()) if not tdf.empty else 0
        t_avg_age = int(tdf["Age (Days)"].mean()) if not tdf.empty else 0

        if "SLA State" in tdf.columns and t_total > 0:
            t_sla = round((tdf["SLA State"] == "Within SLA").sum() / t_total * 100, 1)
        else:
            t_sla = 0

        content = f'<div class="tab-content" data-tab-group="tower" data-tab-id="{safe_id}">'
        content += '<div class="kpi-row">'
        content += _kpi_card("Vulns", t_total, C["accent"])
        content += _kpi_card("Hosts", t_hosts, C["accent"])
        content += _kpi_card("Critical", t_crit, C["critical"])
        content += _kpi_card("High", t_high, C["high"])
        content += _kpi_card("Avg Age", t_avg_age, C["medium"])
        content += _kpi_card("SLA", f"{t_sla}%", C["low"] if t_sla >= 80 else C["critical"])
        content += "</div>"

        content += '<div class="chart-grid">'
        content += _switchable_chart(f"sev-{safe_id}", {
            "pie": chart_severity_pie(tdf, f"{tname} - Severity"),
            "bar": chart_severity_bar(tdf, f"{tname} - Severity"),
        })
        content += f'<div class="chart-card">{_fig_to_div(chart_sla_compliance(tdf, f"{tname} - SLA"))}</div>'
        content += f'<div class="chart-card">{_fig_to_div(chart_age_histogram(tdf, f"{tname} - Age Distribution"))}</div>'
        content += f'<div class="chart-card">{_fig_to_div(chart_exploit_pie(tdf, f"{tname} - Exploits"))}</div>'
        content += f'<div class="chart-card chart-full">{_fig_to_div(chart_top_vulns(tdf, n=10, title=f"{tname} - Top Vulnerabilities"))}</div>'
        content += f'<div class="chart-card chart-full">{_fig_to_div(chart_family_bar(tdf, title=f"{tname} - Plugin Families"))}</div>'
        content += f'<div class="chart-card chart-full">{_fig_to_div(chart_weekly_trend(tdf, tw_info, f"{tname} - Weekly Trend"))}</div>'
        content += f'<div class="chart-card chart-full">{_fig_to_div(chart_aging_buckets(tdf, f"{tname} - Aging Buckets"))}</div>'
        content += "</div>"

        content += f'<div class="section-title">{tname} - Vulnerability Details</div>'
        content += _build_data_table(tdf, f"tbl-{safe_id}")
        content += "</div>"

        tower_tab_contents += content

    tab_buttons += "</div>"

    overview_tab = f'<div class="tab-content active" data-tab-group="tower" data-tab-id="overview">{kpi}{overview_charts}</div>'

    body = tab_buttons + overview_tab + tower_tab_contents

    full_html = _page_wrapper("Vulnerability Operations Center", body, _nav_links("dashboard"))

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(full_html)
    print(f"  Generated: {output_path}")


# ═══════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════

def main():
    print("=" * 60)
    print("  TENABLE VULNERABILITY TRACKER - NOC DASHBOARD GENERATOR")
    print("=" * 60)

    # ── Get API keys ──
    access_key = accessKey
    secret_key = secretKey

    if not access_key or not secret_key:
        print("\n[WARN] No API keys provided. Attempting to load from existing Excel...")
        all_raw = {}
    else:
        # ── Fetch from Tenable ──
        print("\n[1/4] Fetching vulnerability data from Tenable.io...")
        all_raw = fetch_all_tags(access_key, secret_key)

    # ── Aggregate ──
    if all_raw:
        print("\n[2/4] Aggregating data by plugin ID...")
        all_plugins = {}
        for tag_name, records in all_raw.items():
            all_plugins[tag_name] = aggregate_by_plugin(records)

        # ── Update Excel ──
        print("\n[3/4] Updating Excel tracker...")
        update_excel(all_plugins, EXCEL_PATH)
    else:
        print("\n[2/4] Skipping aggregation (no API data).")
        print("[3/4] Skipping Excel update.")

    # ── Load from Excel for dashboard generation ──
    print("\n[4/4] Generating interactive HTML dashboards...")

    if os.path.exists(EXCEL_PATH):
        tag_dfs, weekly_info = load_excel_to_dataframes(EXCEL_PATH)
    else:
        # Build DataFrames from fresh API data if no Excel
        if all_raw:
            tag_dfs = {}
            weekly_info = {}
            for tag_name, records in all_raw.items():
                plugins = aggregate_by_plugin(records)
                tag_dfs[tag_name] = plugins_to_dataframe(plugins, tag_name)
                weekly_info[tag_name] = []
        else:
            print("[ERROR] No Excel file found and no API data fetched. Nothing to generate.")
            sys.exit(1)

    if not tag_dfs:
        print("[ERROR] No data available for dashboard generation.")
        sys.exit(1)

    combined = pd.concat(tag_dfs.values(), ignore_index=True)

    os.makedirs(REPORT_DIR, exist_ok=True)

    # Generate main dashboard
    generate_dashboard_html(tag_dfs, weekly_info, os.path.join(REPORT_DIR, "dashboard.html"))

    # Generate severity breakdowns
    for sev in SEV_ORDER:
        generate_severity_html(combined, sev, os.path.join(REPORT_DIR, f"{sev.lower()}.html"))

    # Generate consolidated
    generate_severity_html(combined, "All", os.path.join(REPORT_DIR, "consolidated.html"))

    print("\n" + "=" * 60)
    print("  COMPLETE! Reports generated in:")
    print(f"  {REPORT_DIR}")
    print("=" * 60)
    print("\nFiles:")
    for f in os.listdir(REPORT_DIR):
        if f.endswith(".html"):
            print(f"  - {f}")
    print(f"\nOpen {os.path.join(REPORT_DIR, 'dashboard.html')} in your browser.")


if __name__ == "__main__":
    main()
